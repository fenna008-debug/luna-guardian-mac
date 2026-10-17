"""Verify bounded document reading and attachment source isolation."""

import sys
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch

from media import allowed_url, receive
from media_worker import extract, process


class MediaTests(unittest.IsolatedAsyncioTestCase):
    def test_urls(self):
        self.assertTrue(allowed_url("https://files.qq.com/test"))
        for url in [
            "http://files.qq.com/a",
            "https://qq.com.evil.test/a",
            "https://127.0.0.1/a",
            "https://user:pass@qq.com/a",
            "file:///etc/passwd",
        ]:
            self.assertFalse(allowed_url(url))

    async def test_local_file_isolation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job = root / "job"
            job.mkdir()
            allowed = root / "temp"
            allowed.mkdir()
            outside = root / "secret.txt"
            outside.write_text("private")
            with self.assertRaises(ValueError):
                await receive(str(outside), "text.txt", job, allowed)
            inside = allowed / "sample.txt"
            inside.write_text("real attachment")
            result = await receive(str(inside), "text.txt", job, allowed)
            self.assertEqual(result.read_text(), "real attachment")

    def test_text_bounds(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.txt"
            path.write_text("文字" * 7000)
            text, scope = extract(path)
            self.assertEqual(len(text), 12000)
            self.assertIn("12000", scope)

    def test_docx_paragraph_and_table(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.docx"
            doc = Document()
            doc.add_paragraph("Test paragraph")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Test table"
            doc.save(path)
            text, _ = extract(path)
            self.assertIn("Test paragraph", text)
            self.assertIn("Test table", text)

    def test_oversized_docx_expansion(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.docx"
            with zipfile.ZipFile(
                path, "w", compression=zipfile.ZIP_DEFLATED
            ) as archive:
                archive.writestr("word/document.xml", b"0" * (51 * 1024 * 1024))
            with self.assertRaises(ValueError):
                extract(path)

    def test_model_failure_returns_real_extracted_text(self):
        fake = types.ModuleType("backend")
        fake.read = lambda name, default: default
        fake.chat = Mock(side_effect=ValueError("provider timeout"))
        with (
            tempfile.TemporaryDirectory() as folder,
            patch.dict(sys.modules, {"backend": fake}),
        ):
            root = Path(folder)
            source = root / "input.txt"
            source.write_text("Sample count: 12. Verify source first.")
            result = process({"action": "file", "path": str(source), "job": str(root)})
            self.assertIn("暂未完成总结", result["answer"])
            self.assertIn("Sample count: 12", Path(result["file"]).read_text())
            self.assertNotIn("provider timeout", result["answer"])

    def test_empty_pdf_not_invented(self):
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            writer.write(path)
            with self.assertRaises(ValueError):
                extract(path)


if __name__ == "__main__":
    unittest.main()
